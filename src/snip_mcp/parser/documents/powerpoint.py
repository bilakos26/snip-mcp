"""PowerPoint document parser — requires python-pptx."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from snip_mcp.parser.documents.base import DocumentFile, DocumentParser, DocumentSection
from snip_mcp.parser.symbols import compute_content_hash


class PowerPointParser(DocumentParser):
    @property
    def supported_extensions(self) -> tuple[str, ...]:
        return (".pptx",)

    @property
    def format_name(self) -> str:
        return "powerpoint"

    def parse(self, file_path: Path, rel_path: str) -> DocumentFile:
        file_bytes = file_path.read_bytes()
        prs = Presentation(file_path)
        sections: list[DocumentSection] = []

        for slide_idx, slide in enumerate(prs.slides, 1):
            title = ""
            content_parts: list[str] = []
            notes = ""

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if shape == slide.shapes.title:
                        title = text
                    elif text:
                        content_parts.append(text)

                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        content_parts.append(row_text)

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()

            if not title:
                title = f"Slide {slide_idx}"

            preview = title
            if content_parts:
                preview += "\n" + "\n".join(content_parts)
            if notes:
                preview += f"\n[Notes: {notes}]"

            sec_id = f"{rel_path}::slide::{title}::{slide_idx}"
            sections.append(
                DocumentSection(
                    id=sec_id,
                    title=title,
                    section_type="slide",
                    file_path=rel_path,
                    line_start=slide_idx,
                    line_end=slide_idx,
                    byte_start=0,
                    byte_end=len(file_bytes),
                    content_preview=preview[:200],
                    level=0,
                    metadata={"slide_number": str(slide_idx), "has_notes": str(bool(notes))},
                )
            )

        return DocumentFile(
            file_path=rel_path,
            format=self.format_name,
            sections=sections,
            content_hash=compute_content_hash(file_bytes.hex()),
            file_size=len(file_bytes),
        )
