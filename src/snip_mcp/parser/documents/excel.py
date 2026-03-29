"""Excel document parser — requires openpyxl."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from snip_mcp.parser.documents.base import DocumentFile, DocumentParser, DocumentSection
from snip_mcp.parser.symbols import compute_content_hash


class ExcelParser(DocumentParser):
    @property
    def supported_extensions(self) -> tuple[str, ...]:
        return (".xlsx",)

    @property
    def format_name(self) -> str:
        return "excel"

    def parse(self, file_path: Path, rel_path: str) -> DocumentFile:
        file_bytes = file_path.read_bytes()
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sections: list[DocumentSection] = []

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(max_row=100, values_only=True))

            # Sheet section
            headers = []
            if rows:
                headers = [str(c) if c is not None else "" for c in rows[0]]

            preview_lines = []
            for row in rows[:6]:
                preview_lines.append(" | ".join(str(c) if c is not None else "" for c in row))
            preview = "\n".join(preview_lines)

            sec_id = f"{rel_path}::sheet::{sheet_name}::{sheet_idx + 1}"
            sections.append(
                DocumentSection(
                    id=sec_id,
                    title=f"Sheet: {sheet_name}",
                    section_type="sheet",
                    file_path=rel_path,
                    line_start=1,
                    line_end=len(rows),
                    byte_start=0,
                    byte_end=len(file_bytes),
                    content_preview=preview[:200],
                    level=0,
                    metadata={
                        "columns": str(len(headers)),
                        "rows": str(len(rows)),
                        "headers": ", ".join(headers[:10]),
                    },
                )
            )

        wb.close()

        return DocumentFile(
            file_path=rel_path,
            format=self.format_name,
            sections=sections,
            content_hash=compute_content_hash(file_bytes.hex()),
            file_size=len(file_bytes),
        )
